"""
Document ingestion: parse files, chunk text, build FTS index.
AI-powered: auto doc-type detection, document summary, per-page summaries.
Hybrid search: metadata → page summaries → chunks.
"""
import os
import re
import json
from database import get_db, put_db

CHUNK_SIZE    = 900
CHUNK_OVERLAP = 200

_PROMPTS_DIR = os.path.join(os.path.dirname(__file__), "prompts")


def _load_prompt(name: str, **kwargs) -> str:
    """
    Load a prompt from prompts/<name>.md.
    - Strips YAML frontmatter (--- ... ---) automatically
    - Substitutes ${variable} and ${variable:default} placeholders
    """
    import re
    path = os.path.join(_PROMPTS_DIR, name + ".md")
    with open(path, "r", encoding="utf-8") as f:
        text = f.read()

    defaults = {}
    if text.startswith("---"):
        end = text.find("\n---", 3)
        if end != -1:
            front = text[3:end]
            in_vars = False
            for line in front.splitlines():
                if line.strip() == "variables:":
                    in_vars = True
                    continue
                if in_vars:
                    if line and not line.startswith(" "):
                        in_vars = False
                        continue
                    m = re.match(r'\s+(\w+):\s*"?(.*?)"?\s*$', line)
                    if m:
                        defaults[m.group(1)] = m.group(2)
            text = text[end + 4:].lstrip("\n")

    values = {**defaults, **{k: str(v) for k, v in kwargs.items()}}

    def _replace(m):
        key = m.group(1)
        inline_default = m.group(2) if m.group(2) is not None else ""
        return values.get(key, inline_default)

    text = re.sub(r'\$\{(\w+)(?::([^}]*))?\}', _replace, text)
    return text

DOC_TYPES = [
    "LPA", "PPM", "Side Letter", "Subscription Agreement",
    "Financial Statements", "Audit Report", "Tax Document",
    "Presentation", "Policy", "Memo", "DDQ", "ESG Report",
    "Fee Schedule", "Distribution Notice", "Capital Call", "Other"
]


# ── File readers ───────────────────────────────────────────────────────────────

def _read_txt(path):
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _read_pdf(path):
    """Returns (full_text, [(page_num, page_text), ...])"""
    try:
        import fitz
        doc = fitz.open(path)
        pages = [(i + 1, page.get_text()) for i, page in enumerate(doc)]
        return "\n\n".join(t for _, t in pages), pages
    except ImportError:
        text = f"[PDF support requires PyMuPDF. File: {os.path.basename(path)}]"
        return text, [(1, text)]
    except Exception as e:
        text = f"[Error reading PDF: {e}]"
        return text, [(1, text)]


def _read_docx(path):
    try:
        from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        return f"[DOCX support requires python-docx. File: {os.path.basename(path)}]"
    except Exception as e:
        return f"[Error reading DOCX: {e}]"


def read_file(path):
    """Returns (content, file_type, pages_or_None)."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        content, pages = _read_pdf(path)
        return content, "pdf", pages
    elif ext in (".docx", ".doc"):
        return _read_docx(path), "docx", None
    elif ext == ".md":
        return _read_txt(path), "markdown", None
    else:
        return _read_txt(path), "text", None


# ── Chunking ───────────────────────────────────────────────────────────────────

def _detect_section(text):
    lines = text.strip().splitlines()
    for line in lines[:3]:
        line = line.strip()
        if re.match(r'^(section|article|clause|schedule|appendix|part)\s+[\d\.]+', line, re.I):
            return line[:80]
        if re.match(r'^[\d]+[\.\)]\s+[A-Z]', line):
            return line[:80]
    return None


def chunk_text_simple(text, size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    text = re.sub(r'\n{3,}', '\n\n', text).strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = start + size
        chunk = text[start:end]
        if end < len(text):
            bp = chunk.rfind('\n\n')
            if bp > size // 3:
                end = start + bp
                chunk = text[start:end]
        if chunk.strip():
            chunks.append({"text": chunk.strip(), "section_ref": _detect_section(chunk)})
        start = end - overlap
        if start >= len(text):
            break
    return chunks


def chunk_by_pages(pages, size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    all_chunks = []
    for page_num, page_text in pages:
        page_chunks = chunk_text_simple(page_text, size, overlap)
        for c in page_chunks:
            c["page_ref"] = f"p.{page_num}"
            all_chunks.append(c)
    return all_chunks


# ── AI Analysis ────────────────────────────────────────────────────────────────

def _get_openai_client():
    try:
        from openai import OpenAI
        key = os.getenv("OPENAI_API_KEY")
        if not key:
            return None
        return OpenAI(api_key=key)
    except Exception:
        return None


def _ai_analyze_document(text_sample, filename):
    """
    Use LLM to detect document type, generate summary, and key topics.
    Returns dict: {ai_doc_type, summary, key_topics}
    """
    client = _get_openai_client()
    if not client:
        return {"ai_doc_type": "Other", "summary": "", "key_topics": []}

    sample = text_sample[:4000].strip()
    types_str = ", ".join(DOC_TYPES)

    prompt = _load_prompt("doc_analyze",
        filename=filename,
        sample=sample,
        doc_types=types_str,
    )

    try:
        resp = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{"role": "user", "content": prompt}],
            response_format={"type": "json_object"},
            max_tokens=600,
            temperature=0,
        )
        data = json.loads(resp.choices[0].message.content)
        return {
            "ai_doc_type": data.get("doc_type", "Other"),
            "summary": data.get("summary", ""),
            "key_topics": data.get("key_topics", []),
        }
    except Exception:
        return {"ai_doc_type": "Other", "summary": "", "key_topics": []}


def _ai_summarize_pages(pages, doc_name):
    """
    Generate a short summary for each page (batched in groups of 5).
    Returns list of {page_num, summary}.
    """
    client = _get_openai_client()
    if not client:
        return []

    # Limit to first 30 pages to keep costs low
    pages_to_summarize = [p for p in pages if p[1].strip()][:30]
    if not pages_to_summarize:
        return []

    results = []
    # Batch 5 pages per LLM call
    batch_size = 5
    for i in range(0, len(pages_to_summarize), batch_size):
        batch = pages_to_summarize[i:i + batch_size]
        pages_text = ""
        for page_num, page_text in batch:
            pages_text += f"\n\n--- PAGE {page_num} ---\n{page_text[:800]}"

        prompt = _load_prompt("doc_page_summary",
            doc_name=doc_name,
            pages_text=pages_text,
        )

        try:
            resp = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[{"role": "user", "content": prompt}],
                response_format={"type": "json_object"},
                max_tokens=800,
                temperature=0,
            )
            data = json.loads(resp.choices[0].message.content)
            for p in data.get("pages", []):
                if p.get("page_num") and p.get("summary"):
                    results.append({"page_num": p["page_num"], "summary": p["summary"]})
        except Exception:
            # On failure, add empty summaries for this batch
            for page_num, _ in batch:
                results.append({"page_num": page_num, "summary": ""})

    return results


# ── Ingestion ──────────────────────────────────────────────────────────────────

def ingest_fund_document(filepath, name, doc_type, uploaded_by, run_ai=True):
    """
    Parse file, AI-analyze, store in DB with summary + page summaries, build FTS.
    Returns fund_documents.id
    """
    content, file_type, pages = read_file(filepath)
    page_count = len(pages) if pages else 0

    # AI analysis
    ai_meta = {"ai_doc_type": doc_type, "summary": "", "key_topics": []}
    page_summaries = []
    if run_ai and content.strip():
        ai_meta = _ai_analyze_document(content, name)
        # Use AI doc type only if user passed "Other" or default
        if doc_type in ("Other", "") or not doc_type:
            doc_type = ai_meta["ai_doc_type"]
        if pages:
            page_summaries = _ai_summarize_pages(pages, name)

    # Chunk
    if pages:
        chunks = chunk_by_pages(pages)
    else:
        chunks = chunk_text_simple(content)
        for c in chunks:
            c["page_ref"] = None

    conn = get_db()
    cur = conn.cursor()

    cur.execute("""
        INSERT INTO fund_documents
            (name, doc_type, ai_doc_type, filepath, content, summary, key_topics,
             page_count, uploaded_by)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
        RETURNING id
    """, (
        name, doc_type, ai_meta["ai_doc_type"],
        filepath, content[:10000],
        ai_meta["summary"],
        json.dumps(ai_meta["key_topics"]),
        page_count, uploaded_by,
    ))
    doc_id = cur.fetchone()[0]

    # Update search_meta tsvector (name + summary + key_topics)
    meta_text = f"{name} {ai_meta['summary']} {' '.join(ai_meta['key_topics'])}"
    cur.execute("""
        UPDATE fund_documents SET search_meta = to_tsvector('english', %s) WHERE id=%s
    """, (meta_text, doc_id))

    # Insert chunks
    for i, chunk in enumerate(chunks):
        cur.execute("""
            INSERT INTO document_chunks
                (document_id, chunk_text, chunk_index, page_ref, section_ref)
            VALUES (%s, %s, %s, %s, %s)
        """, (doc_id, chunk["text"], i,
              chunk.get("page_ref"), chunk.get("section_ref")))

    # Insert page summaries
    for ps in page_summaries:
        if ps["summary"]:
            cur.execute("""
                INSERT INTO page_summaries (document_id, page_num, summary, search_vector)
                VALUES (%s, %s, %s, to_tsvector('english', %s))
            """, (doc_id, ps["page_num"], ps["summary"], ps["summary"]))

    conn.commit()
    put_db(conn)
    return doc_id


def ingest_investor_document(filepath, name, doc_type, investor_session_id):
    """Ingest an investor-specific document. Returns investor_documents.id"""
    content, file_type, pages = read_file(filepath)

    if pages:
        chunks = chunk_by_pages(pages)
    else:
        chunks = chunk_text_simple(content)
        for c in chunks:
            c["page_ref"] = None

    conn = get_db()
    cur = conn.cursor()

    cur.execute("""
        INSERT INTO investor_documents
            (investor_session_id, name, filepath, doc_type, content)
        VALUES (%s, %s, %s, %s, %s)
        RETURNING id
    """, (investor_session_id, name, filepath, doc_type, content[:10000]))
    inv_doc_id = cur.fetchone()[0]

    for i, chunk in enumerate(chunks):
        cur.execute("""
            INSERT INTO investor_doc_chunks (investor_doc_id, chunk_text, chunk_index)
            VALUES (%s, %s, %s)
        """, (inv_doc_id, chunk["text"], i))

    conn.commit()
    put_db(conn)
    return inv_doc_id


def ingest_from_folder(folder_path, uploaded_by=1):
    """Bulk ingest all supported files from a folder."""
    supported = {".txt", ".md", ".pdf", ".docx", ".doc"}
    results = []
    for fname in os.listdir(folder_path):
        ext = os.path.splitext(fname)[1].lower()
        if ext not in supported:
            continue
        filepath = os.path.join(folder_path, fname)
        try:
            doc_id = ingest_fund_document(filepath, fname, "Other", uploaded_by, run_ai=True)
            results.append({"file": fname, "doc_id": doc_id, "status": "ok"})
        except Exception as e:
            results.append({"file": fname, "status": "error", "error": str(e)})
    return results
